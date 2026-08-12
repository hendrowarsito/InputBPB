"""Uji modul :mod:`btb_io`.

Jalankan dengan ``pytest test_btb_io.py`` atau ``python test_btb_io.py``.

Uji yang membutuhkan PDF contoh akan dilewati bila berkas
``samples/btb_print_kota_jakarta_2026.pdf`` tidak tersedia, atau bila
``tesseract`` belum terpasang di sistem.
"""

from __future__ import annotations

import io
import os

import pandas as pd
import pytest

import btb_io as btb

SAMPLE_PDF = os.path.join(os.path.dirname(__file__), "samples", "btb_print_kota_jakarta_2026.pdf")

# Beberapa angka acuan dari lembar contoh (KOTA JAKARTA, 2026).
EXPECTED = {
    ("Pondasi", "Mewah"): 941_613,
    ("Struktur", "Tinggi (High-Rise) Grade A"): 3_380_573,
    ("Persiapan", "Sederhana"): 0,
    ("Pintu dan Jendela", "Model Mall (Grade B)"): 2_436,
    ("Utilitas", "Model Hotel (Bintang 4)"): 3_072_454,
    (btb.TOTAL_A, "Mewah"): 7_170_995,
    (btb.TOTAL_B, "Model Apartemen (Grade B)"): 1_181_497,
    (btb.PPN_ROW, "Menengah"): 600_005,
    (btb.ROUNDED, "Model Hotel (Bintang 4)"): 12_920_000,
}


def has_ocr() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


needs_sample = pytest.mark.skipif(
    not os.path.exists(SAMPLE_PDF), reason="PDF contoh tidak tersedia"
)
needs_ocr = pytest.mark.skipif(not has_ocr(), reason="tesseract belum terpasang")


# ---------------------------------------------------------------------------
# Penguraian angka
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1,780,811", 1_780_811),   # ribuan gaya Inggris (format dokumen MAPPI)
        ("1.780.811", 1_780_811),   # ribuan gaya Indonesia
        ("1,234.56", 1234.56),
        ("1.234,56", 1234.56),
        ("12,5", 12.5),
        ("0", 0),
        ("Rp 941,613", 941_613),
        ("-1,000", -1000),
        ("", None),
        ("-", None),
        (None, None),
        (2026, 2026),
    ],
)
def test_parse_number(raw, expected):
    assert btb.parse_number(raw) == expected


def test_match_row_label_toleran_terhadap_salah_ocr():
    assert btb.match_row_label("TOTAL BIAYA LANGSUNG ( A )") == btb.TOTAL_A
    assert btb.match_row_label("Keuntungan  Kontraktor") == "Keuntungan Kontraktor"
    assert btb.match_row_label("Pintu dan Jendela") == "Pintu dan Jendela"
    assert btb.match_row_label("catatan kaki halaman") is None


def test_match_building_column():
    assert btb.match_building_column("MEWAH 2 LANTAI Rp./m2") == "Mewah"
    assert btb.match_building_column("TINGGI (HIGH RISE) GRADE A") == "Tinggi (High-Rise) Grade A"
    assert btb.match_building_column("MODEL HOTEL (BINTANG 4)") == "Model Hotel (Bintang 4)"


# ---------------------------------------------------------------------------
# Pembacaan PDF
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def sample_sheet():
    return btb.read_pdf(SAMPLE_PDF, os.path.basename(SAMPLE_PDF))


@needs_sample
@needs_ocr
def test_metadata_terbaca(sample_sheet):
    assert sample_sheet.provinsi == "PROV. DKI JAKARTA"
    assert sample_sheet.kota_kabupaten == "KOTA JAKARTA"
    assert sample_sheet.tahun == 2026


@needs_sample
@needs_ocr
def test_seluruh_sel_terisi(sample_sheet):
    values = sample_sheet.values.drop(index=[label for label in btb.ROW_LABELS if btb.ROW_KIND[label] == "section"])
    assert values.shape == (19, 11)
    assert not values.isna().any().any(), "masih ada sel kosong"


@needs_sample
@needs_ocr
@pytest.mark.parametrize("key,expected", list(EXPECTED.items()))
def test_nilai_acuan(sample_sheet, key, expected):
    elemen, kolom = key
    assert float(sample_sheet.values.loc[elemen, kolom]) == expected


@needs_sample
@needs_ocr
def test_validasi_aritmatika_lolos(sample_sheet):
    """Total, PPN, dan pembulatan harus konsisten - bukti angka terbaca benar."""
    assert btb.validate(sample_sheet) == []


# ---------------------------------------------------------------------------
# Penulisan dan pembacaan XLSX
# ---------------------------------------------------------------------------


def dummy_sheet(kota: str = "KOTA UJI", tahun: int = 2026) -> btb.BTBSheet:
    sheet = btb.BTBSheet(provinsi="PROV. UJI", kota_kabupaten=kota, tahun=tahun)
    for i, label in enumerate(btb.ROW_LABELS):
        if btb.ROW_KIND[label] == "section":
            continue
        for j, column in enumerate(btb.BUILDING_COLUMNS):
            sheet.values.loc[label, column] = float((i + 1) * 1000 + j)
    return btb.recompute_totals(sheet)


def test_recompute_totals_konsisten():
    assert btb.validate(dummy_sheet()) == []


def test_simpan_dan_baca_ulang(tmp_path):
    path = str(tmp_path / "btb_data.xlsx")
    sheet = dummy_sheet()

    data = btb.save_sheet(path, sheet)
    assert len(data) == len(btb.ROW_LABELS)

    # menyimpan ulang kota/tahun yang sama harus menimpa, bukan menggandakan
    data = btb.save_sheet(path, sheet)
    assert len(data) == len(btb.ROW_LABELS)

    data = btb.save_sheet(path, dummy_sheet("KOTA LAIN", 2025))
    assert len(data) == 2 * len(btb.ROW_LABELS)
    assert len(btb.summary(data)) == 2

    with open(path, "rb") as handle:
        kembali = btb.read_spreadsheet(handle, "btb_data.xlsx")
    assert len(kembali) == 2
    terpilih = next(s for s in kembali if s.kota_kabupaten == "KOTA UJI")
    assert terpilih.tahun == 2026
    pd.testing.assert_frame_equal(
        terpilih.values.astype(float), sheet.values.astype(float), check_names=False
    )


def test_xlsx_menyimpan_angka_sebagai_numerik(tmp_path):
    from openpyxl import load_workbook

    path = str(tmp_path / "btb_data.xlsx")
    btb.save_sheet(path, dummy_sheet())
    workbook = load_workbook(path)
    assert btb.SHEET_NAME in workbook.sheetnames
    worksheet = workbook[btb.SHEET_NAME]
    header = [cell.value for cell in worksheet[1]]
    assert header == btb.ID_COLUMNS + btb.BUILDING_COLUMNS

    kolom_nilai = header.index("Mewah") + 1
    baris_pondasi = next(
        row for row in range(2, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=header.index("Elemen") + 1).value == "Pondasi"
    )
    cell = worksheet.cell(row=baris_pondasi, column=kolom_nilai)
    assert isinstance(cell.value, (int, float)), "angka harus tersimpan numerik, bukan teks"
    assert cell.number_format == "#,##0"


def test_baca_csv_bentuk_panjang(tmp_path):
    path = str(tmp_path / "btb.csv")
    btb.wide_to_long(dummy_sheet().to_wide()).to_csv(path, index=False)
    with open(path, "rb") as handle:
        sheets = btb.read_spreadsheet(handle, "btb.csv")
    assert len(sheets) == 1
    assert sheets[0].kota_kabupaten == "KOTA UJI"
    assert btb.validate(sheets[0]) == []


def test_unduhan_wide_dan_long():
    wide = dummy_sheet().to_wide()
    for layout, baris in (("wide", len(btb.ROW_LABELS)), ("long", len(btb.ROW_LABELS) * 11)):
        buffer = btb.build_download(wide, layout)
        assert isinstance(buffer, io.BytesIO)
        assert len(pd.read_excel(buffer)) == baris


def test_berkas_tanpa_kolom_elemen_ditolak(tmp_path):
    path = str(tmp_path / "salah.csv")
    pd.DataFrame({"A": [1], "B": [2]}).to_csv(path, index=False)
    with pytest.raises(btb.BTBError):
        btb.read_spreadsheet(open(path, "rb"), "salah.csv")


def test_load_workbook_data_berkas_belum_ada(tmp_path):
    kosong = btb.load_workbook_data(str(tmp_path / "belum-ada.xlsx"))
    assert kosong.empty
    assert list(kosong.columns) == btb.ID_COLUMNS + btb.BUILDING_COLUMNS


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
