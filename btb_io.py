"""Pembacaan dan penulisan data BTB (Biaya Teknis Bangunan) MAPPI.

Modul ini murni logika (tanpa Streamlit) sehingga bisa diuji terpisah.

Dua hal yang ditangani di sini:

1. **Membaca** lembar BTB dari PDF cetakan portal MAPPI
   (``member.mappi.or.id/btb_print-*.html`` yang di-print ke PDF) maupun dari
   berkas XLSX/CSV.

   Cetakan PDF dari portal MAPPI tidak memiliki lapisan teks sama sekali:
   seluruh huruf digambar sebagai kurva vektor (0 karakter, 0 font di dalam
   berkas). Karena itu ``pdfplumber.extract_tables()`` selalu mengembalikan
   daftar kosong dan kode lama menghasilkan tabel kosong. Solusinya: garis
   tabel tetap ada sebagai objek vektor, jadi grid sel direkonstruksi dari
   garis tersebut lalu setiap sel dibaca satu per satu — dari lapisan teks bila
   ada, atau melalui OCR bila tidak ada.

2. **Menulis** hasilnya ke XLSX yang rapi: satu sheet induk ``BTB Data``
   berformat lebar (mengikuti bentuk dokumen aslinya), angka bernilai numerik
   asli dengan format ribuan, dan penyimpanan idempoten (menyimpan ulang
   kota/tahun yang sama menimpa, bukan menggandakan).
"""

from __future__ import annotations

import difflib
import io
import os
import re
from dataclasses import dataclass, field
from typing import Any, Sequence

import pandas as pd

# ---------------------------------------------------------------------------
# Struktur baku lembar BTB
# ---------------------------------------------------------------------------

SHEET_NAME = "BTB Data"

#: Kolom identitas yang selalu mendahului kolom tipe bangunan.
ID_COLUMNS = ["Provinsi", "Kota/Kabupaten", "Tahun", "No", "Kelompok", "Elemen"]

#: 11 tipe bangunan pada formulir BTB, berurutan sesuai dokumen.
COLUMN_SPECS: list[tuple[str, str, str, tuple[str, ...]]] = [
    # (kelompok, nama kolom baku, spesifikasi, kata kunci pencocokan)
    ("Bangunan Rumah Tinggal", "Mewah", "2 Lantai", ("mewah",)),
    ("Bangunan Rumah Tinggal", "Menengah", "2 Lantai", ("menengah",)),
    ("Bangunan Rumah Tinggal", "Sederhana", "1 Lantai", ("sederhana",)),
    ("Bangunan Perkebunan", "Semi Permanen", "1 Lantai", ("perkebunan", "permanen")),
    ("Bangunan Gudang", "Gudang", "1 Lantai", ("gudang",)),
    ("Bangunan Gedung Bertingkat", "Rendah (Low-Rise)", "3 Lantai (< 5 Lantai)", ("rendah", "low")),
    ("Bangunan Gedung Bertingkat", "Sedang (Mid-Rise)", "8 Lantai + 1 Basement (5 - 8 Lantai)", ("sedang", "mid")),
    ("Bangunan Gedung Bertingkat", "Tinggi (High-Rise) Grade A", "32 Lantai + 2 Basement", ("tinggi", "high")),
    ("Model Mall", "Model Mall (Grade B)", "4 Lantai + 1 Basement", ("mall",)),
    ("Model Hotel", "Model Hotel (Bintang 4)", "24 Lantai + 2 Basement", ("hotel", "bintang")),
    ("Model Apartemen", "Model Apartemen (Grade B)", "18 Lantai + 1 Basement", ("apartemen", "apartment")),
]

#: Nama kolom tipe bangunan sebagaimana ditulis ke XLSX.
BUILDING_COLUMNS = [spec[1] for spec in COLUMN_SPECS]

# Jenis baris:
#   section  -> judul kelompok, tanpa nilai
#   item     -> komponen biaya yang dijumlahkan
#   subtotal -> hasil penjumlahan item di kelompoknya
#   derived  -> turunan dari subtotal (A+B, PPN, dst.)
ROW_SPECS: list[tuple[str, str, str]] = [
    ("A. BIAYA LANGSUNG", "A. BIAYA LANGSUNG", "section"),
    ("A. BIAYA LANGSUNG", "Persiapan", "item"),
    ("A. BIAYA LANGSUNG", "Pondasi", "item"),
    ("A. BIAYA LANGSUNG", "Struktur", "item"),
    ("A. BIAYA LANGSUNG", "Rangka Atap", "item"),
    ("A. BIAYA LANGSUNG", "Penutup Atap", "item"),
    ("A. BIAYA LANGSUNG", "Plafon", "item"),
    ("A. BIAYA LANGSUNG", "Dinding", "item"),
    ("A. BIAYA LANGSUNG", "Pintu dan Jendela", "item"),
    ("A. BIAYA LANGSUNG", "Lantai", "item"),
    ("A. BIAYA LANGSUNG", "Utilitas", "item"),
    ("A. BIAYA LANGSUNG", "TOTAL BIAYA LANGSUNG ( A )", "subtotal"),
    ("B. BIAYA TIDAK LANGSUNG", "B. BIAYA TIDAK LANGSUNG", "section"),
    ("B. BIAYA TIDAK LANGSUNG", "Professional Fee", "item"),
    ("B. BIAYA TIDAK LANGSUNG", "Biaya Perijinan", "item"),
    ("B. BIAYA TIDAK LANGSUNG", "Keuntungan Kontraktor", "item"),
    ("B. BIAYA TIDAK LANGSUNG", "TOTAL BIAYA TIDAK LANGSUNG ( B )", "subtotal"),
    ("REKAPITULASI", "TOTAL BIAYA PEMBANGUNAN BARU ( A + B )", "derived"),
    ("REKAPITULASI", "PPN 11%", "derived"),
    ("REKAPITULASI", "TOTAL BIAYA PEMB. BARU SETELAH PPN", "derived"),
    ("REKAPITULASI", "PEMBULATAN", "derived"),
]

ROW_LABELS = [spec[1] for spec in ROW_SPECS]
ROW_KIND = {label: kind for _, label, kind in ROW_SPECS}
ROW_GROUP = {label: group for group, label, _ in ROW_SPECS}

TOTAL_A = "TOTAL BIAYA LANGSUNG ( A )"
TOTAL_B = "TOTAL BIAYA TIDAK LANGSUNG ( B )"
TOTAL_AB = "TOTAL BIAYA PEMBANGUNAN BARU ( A + B )"
PPN_ROW = "PPN 11%"
TOTAL_PPN = "TOTAL BIAYA PEMB. BARU SETELAH PPN"
ROUNDED = "PEMBULATAN"

ITEMS_A = [label for group, label, kind in ROW_SPECS if kind == "item" and group.startswith("A.")]
ITEMS_B = [label for group, label, kind in ROW_SPECS if kind == "item" and group.startswith("B.")]

DEFAULT_PPN_RATE = 0.11


class BTBError(Exception):
    """Kesalahan yang layak ditampilkan apa adanya kepada pengguna."""


# ---------------------------------------------------------------------------
# Hasil pembacaan
# ---------------------------------------------------------------------------


@dataclass
class BTBSheet:
    """Satu lembar BTB: metadata + matriks elemen biaya x tipe bangunan."""

    provinsi: str = ""
    kota_kabupaten: str = ""
    tahun: int | None = None
    #: index = elemen (ROW_LABELS), columns = BUILDING_COLUMNS, nilai Rp/m2.
    values: pd.DataFrame = field(default_factory=lambda: empty_matrix())
    #: catatan proses baca (mis. kolom hasil OCR yang meragukan).
    notes: list[str] = field(default_factory=list)
    source: str = ""

    def is_complete(self) -> bool:
        return bool(self.kota_kabupaten) and self.tahun is not None

    def to_wide(self) -> pd.DataFrame:
        """Bentuk lebar: satu baris per elemen, 11 kolom tipe bangunan."""
        df = self.values.reindex(ROW_LABELS).copy()
        df.insert(0, "Elemen", df.index)
        df.insert(0, "Kelompok", [ROW_GROUP[label] for label in df.index])
        df.insert(0, "No", range(1, len(df) + 1))
        df.insert(0, "Tahun", self.tahun)
        df.insert(0, "Kota/Kabupaten", self.kota_kabupaten)
        df.insert(0, "Provinsi", self.provinsi)
        return df.reset_index(drop=True)

    def to_long(self) -> pd.DataFrame:
        """Bentuk panjang (tidy) untuk analisis/pivot."""
        wide = self.to_wide()
        long = wide.melt(
            id_vars=ID_COLUMNS,
            value_vars=BUILDING_COLUMNS,
            var_name="Tipe Bangunan",
            value_name="Nilai (Rp/m2)",
        )
        spec = {name: (grp, sp) for grp, name, sp, _ in COLUMN_SPECS}
        long["Kelompok Bangunan"] = long["Tipe Bangunan"].map(lambda c: spec.get(c, ("", ""))[0])
        long["Spesifikasi"] = long["Tipe Bangunan"].map(lambda c: spec.get(c, ("", ""))[1])
        long = long[
            ID_COLUMNS
            + ["Kelompok Bangunan", "Tipe Bangunan", "Spesifikasi", "Nilai (Rp/m2)"]
        ]
        return long.sort_values(["No", "Tipe Bangunan"], kind="stable").reset_index(drop=True)


def empty_matrix() -> pd.DataFrame:
    return pd.DataFrame(pd.NA, index=list(ROW_LABELS), columns=list(BUILDING_COLUMNS), dtype="Float64")


# ---------------------------------------------------------------------------
# Utilitas angka & teks
# ---------------------------------------------------------------------------

_NUMBER_JUNK = re.compile(r"[^0-9,.\-]")


def parse_number(raw: Any) -> float | None:
    """Ubah teks angka BTB menjadi float.

    Menerima format ribuan Inggris (``1,780,811``) maupun Indonesia
    (``1.780.811``), termasuk desimal (``1.234,56`` / ``1,234.56``).
    Mengembalikan ``None`` bila sel benar-benar kosong.
    """
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return None if pd.isna(raw) else float(raw)
    text = _NUMBER_JUNK.sub("", str(raw).strip())
    text = text.replace("--", "-")
    if text in {"", "-", ".", ","}:
        return None

    negative = text.startswith("-")
    text = text.lstrip("-")

    has_comma, has_dot = "," in text, "." in text
    if has_comma and has_dot:
        # pemisah desimal = tanda baca yang muncul paling akhir
        dec = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if dec == "," else ","
        text = text.replace(thousands, "").replace(dec, ".")
    elif has_comma or has_dot:
        sep = "," if has_comma else "."
        parts = text.split(sep)
        tail = parts[-1]
        # ``1,234`` / ``1.234.567`` -> pemisah ribuan; ``12,5`` -> desimal
        if len(parts) > 2 or (len(tail) == 3 and len(parts[0]) <= 3 and parts[0] != ""):
            text = text.replace(sep, "")
        else:
            text = text.replace(sep, ".")
    if text in {"", "."}:
        return None
    try:
        value = float(text)
    except ValueError:
        return None
    return -value if negative else value


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def match_row_label(text: str) -> str | None:
    """Cocokkan teks hasil OCR ke label baris baku."""
    norm = _norm(text)
    if not norm:
        return None
    table = {_norm(label): label for label in ROW_LABELS}
    if norm in table:
        return table[norm]
    hit = difflib.get_close_matches(norm, list(table), n=1, cutoff=0.72)
    return table[hit[0]] if hit else None


def match_building_column(text: str) -> str | None:
    """Cocokkan teks kepala kolom ke nama tipe bangunan baku."""
    norm = _norm(text)
    if not norm:
        return None
    best, best_score = None, 0
    for _, name, _, keywords in COLUMN_SPECS:
        score = sum(1 for kw in keywords if kw in norm)
        if score > best_score:
            best, best_score = name, score
    return best


# ---------------------------------------------------------------------------
# Rekonstruksi grid tabel dari garis vektor PDF
# ---------------------------------------------------------------------------


def _cluster(values: Sequence[float], tol: float = 2.0) -> list[float]:
    values = sorted(values)
    if not values:
        return []
    groups: list[list[float]] = [[values[0]]]
    for value in values[1:]:
        if value - groups[-1][-1] <= tol:
            groups[-1].append(value)
        else:
            groups.append([value])
    return [sum(g) / len(g) for g in groups]


def detect_grid(page) -> tuple[list[float], list[float]]:
    """Kembalikan (batas kolom x, batas baris y) tabel utama pada halaman.

    Garis tabel pada cetakan MAPPI digambar sebagai persegi/kurva sangat tipis,
    bukan objek ``line``; ketiganya diperiksa.
    """
    segments = list(page.rects) + list(page.curves) + list(page.lines)
    vertical, horizontal = [], []
    for seg in segments:
        width = seg["x1"] - seg["x0"]
        height = seg["bottom"] - seg["top"]
        if width < 2.5 and height > 6:
            vertical.append(seg)
        if height < 2.5 and width > 6:
            horizontal.append(seg)
    if len(vertical) < 4 or len(horizontal) < 4:
        return [], []

    xs = _cluster([(s["x0"] + s["x1"]) / 2 for s in vertical])
    ys = _cluster([(s["top"] + s["bottom"]) / 2 for s in horizontal])

    # Panjang total garis pada tiap klaster: pemisah kolom/baris sejati
    # membentang hampir sepanjang tabel, garis kepala tabel hanya sebagian.
    cover_x: dict[float, float] = {}
    for seg in vertical:
        key = min(xs, key=lambda x: abs(x - (seg["x0"] + seg["x1"]) / 2))
        cover_x[key] = cover_x.get(key, 0.0) + seg["bottom"] - seg["top"]
    cover_y: dict[float, float] = {}
    for seg in horizontal:
        key = min(ys, key=lambda y: abs(y - (seg["top"] + seg["bottom"]) / 2))
        cover_y[key] = cover_y.get(key, 0.0) + seg["x1"] - seg["x0"]

    max_x, max_y = max(cover_x.values()), max(cover_y.values())
    cols = sorted(x for x, cov in cover_x.items() if cov > 0.6 * max_x)
    rows = sorted(y for y, cov in cover_y.items() if cov > 0.9 * max_y)
    return cols, rows


# ---------------------------------------------------------------------------
# Pembacaan PDF
# ---------------------------------------------------------------------------


def _page_has_text(page) -> bool:
    return len(page.chars) > 0


def _text_cell(page, x0, top, x1, bottom) -> str:
    """Ambil teks lapisan PDF di dalam satu sel."""
    words = [
        w
        for w in page.extract_words(use_text_flow=False, keep_blank_chars=False)
        if x0 <= (w["x0"] + w["x1"]) / 2 <= x1 and top <= (w["top"] + w["bottom"]) / 2 <= bottom
    ]
    words.sort(key=lambda w: (round(w["top"], 1), w["x0"]))
    return " ".join(w["text"] for w in words).strip()


class _OCR:
    """Pembungkus tesseract: dimuat sekali, memberi pesan jelas bila absen."""

    def __init__(self, pdf_bytes: bytes, scale: int = 4):
        try:
            import pypdfium2 as pdfium
            import pytesseract
            from PIL import ImageOps
        except ImportError as exc:  # pragma: no cover - tergantung lingkungan
            raise BTBError(
                "PDF ini tidak memiliki lapisan teks sehingga perlu OCR, "
                "tetapi paket OCR belum terpasang. Jalankan "
                "`pip install -r requirements.txt`, atau unggah data dalam "
                "format XLSX/CSV."
            ) from exc
        try:
            pytesseract.get_tesseract_version()
        except Exception as exc:  # pragma: no cover - tergantung lingkungan
            raise BTBError(
                "PDF ini tidak memiliki lapisan teks sehingga perlu OCR, "
                "tetapi program `tesseract` tidak ditemukan di sistem "
                "(di Streamlit Cloud: pastikan `packages.txt` memuat "
                "`tesseract-ocr`). Alternatifnya, unggah data dalam format "
                "XLSX/CSV."
            ) from exc

        self._pytesseract = pytesseract
        self._ImageOps = ImageOps
        self.scale = scale
        self._doc = pdfium.PdfDocument(io.BytesIO(pdf_bytes))
        self._pages: dict[int, Any] = {}
        langs = set(pytesseract.get_languages(config="") or [])
        self.lang = "ind+eng" if "ind" in langs else "eng"

    def _image(self, page_index: int):
        if page_index not in self._pages:
            self._pages[page_index] = self._doc[page_index].render(scale=self.scale).to_pil().convert("L")
        return self._pages[page_index]

    def _crop(self, page_index: int, x0, top, x1, bottom, pad: int = 2):
        image = self._image(page_index)
        box = (
            max(0, int(x0 * self.scale) + pad),
            max(0, int(top * self.scale) + pad),
            min(image.width, int(x1 * self.scale) - pad),
            min(image.height, int(bottom * self.scale) - pad),
        )
        if box[2] <= box[0] or box[3] <= box[1]:
            return None
        crop = image.crop(box)
        # Baris abu-abu pada dokumen (PEMBULATAN, TOTAL A+B) nyaris hilang bila
        # tidak dinaikkan kontrasnya; margin putih membantu segmentasi tesseract.
        crop = self._ImageOps.autocontrast(crop, cutoff=1)
        return self._ImageOps.expand(crop, border=12, fill=255)

    def text(self, page_index: int, x0, top, x1, bottom, numeric: bool = False) -> str:
        crop = self._crop(page_index, x0, top, x1, bottom)
        if crop is None:
            return ""
        if numeric:
            configs = [
                "--psm 7 -c tessedit_char_whitelist=0123456789,.-",
                "--psm 6 -c tessedit_char_whitelist=0123456789,.-",
                "--psm 10 -c tessedit_char_whitelist=0123456789,.-",
            ]
            lang = "eng"
        else:
            configs = ["--psm 6", "--psm 7"]
            lang = self.lang
        for config in configs:
            out = self._pytesseract.image_to_string(crop, lang=lang, config=config).strip()
            out = " ".join(out.split())
            if out:
                return out
        return ""


_META_PATTERNS = {
    "provinsi": re.compile(r"provinsi\s*[:.]?\s*(.+)", re.I),
    "kota": re.compile(r"kota\s*/?\s*kab(?:upaten)?\s*[:.]?\s*(.+)", re.I),
    "tahun": re.compile(r"tahun\s*[:.]?\s*(\d{4})", re.I),
}


def _parse_metadata(text: str, sheet: BTBSheet) -> None:
    for line in text.splitlines():
        line = " ".join(line.split())
        if not line:
            continue
        if not sheet.provinsi:
            m = _META_PATTERNS["provinsi"].search(line)
            if m:
                sheet.provinsi = m.group(1).strip(" :.-")
        if not sheet.kota_kabupaten:
            m = _META_PATTERNS["kota"].search(line)
            if m:
                sheet.kota_kabupaten = m.group(1).strip(" :.-")
        if sheet.tahun is None:
            m = _META_PATTERNS["tahun"].search(line)
            if m:
                sheet.tahun = int(m.group(1))


def read_pdf(file: Any, filename: str = "") -> BTBSheet:
    """Baca lembar BTB dari PDF cetakan portal MAPPI.

    Memakai lapisan teks bila tersedia, dan jatuh ke OCR bila PDF hanya berisi
    kurva vektor (kasus cetakan ``btb_print-*.html``).
    """
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise BTBError("Paket `pdfplumber` belum terpasang.") from exc

    data = file.read() if hasattr(file, "read") else open(file, "rb").read()
    if hasattr(file, "seek"):
        try:
            file.seek(0)
        except Exception:
            pass

    sheet = BTBSheet(source=filename or getattr(file, "name", "") or "PDF")
    ocr: _OCR | None = None
    found_rows = 0

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_index, page in enumerate(pdf.pages):
            cols, rows = detect_grid(page)
            if len(cols) < 3 or len(rows) < 3:
                continue

            use_ocr = not _page_has_text(page)
            if use_ocr and ocr is None:
                ocr = _OCR(data)

            def cell(x0, top, x1, bottom, numeric=False) -> str:
                if use_ocr:
                    return ocr.text(page_index, x0, top, x1, bottom, numeric=numeric)
                return _text_cell(page, x0, top, x1, bottom)

            # --- metadata di atas tabel ---
            if not sheet.is_complete() and rows[0] > 20:
                meta_text = (
                    ocr.text(page_index, 0, 0, page.width, rows[0] - 1)
                    if use_ocr
                    else (page.crop((0, 0, page.width, rows[0] - 1)).extract_text() or "")
                )
                # OCR satu blok mengembalikan satu baris panjang; pisahkan per label
                meta_text = re.sub(r"\s(Kota|Tahun|Provinsi)\b", r"\n\1", meta_text)
                _parse_metadata(meta_text, sheet)

            # --- kepala kolom (blok pertama) untuk verifikasi urutan ---
            header_names: list[str | None] = []
            for j in range(1, len(cols) - 1):
                header_names.append(match_building_column(cell(cols[j], rows[0], cols[j + 1], rows[1])))

            n_values = len(cols) - 2
            if n_values != len(BUILDING_COLUMNS):
                sheet.notes.append(
                    f"Halaman {page_index + 1}: terbaca {n_values} kolom tipe bangunan, "
                    f"seharusnya {len(BUILDING_COLUMNS)}. Periksa hasilnya sebelum menyimpan."
                )
            for pos, name in enumerate(header_names[: len(BUILDING_COLUMNS)]):
                expected = BUILDING_COLUMNS[pos]
                if name and name != expected:
                    sheet.notes.append(
                        f"Kepala kolom ke-{pos + 1} terbaca sebagai '{name}', "
                        f"dipetakan ke urutan baku '{expected}'."
                    )

            # --- baris data ---
            for i in range(1, len(rows) - 1):
                top, bottom = rows[i], rows[i + 1]
                if bottom - top < 5:
                    continue
                label_raw = cell(cols[0], top, cols[1], bottom)
                label = match_row_label(label_raw)
                if label is None:
                    if label_raw:
                        sheet.notes.append(f"Baris '{label_raw}' tidak dikenali, dilewati.")
                    continue
                found_rows += 1
                if ROW_KIND[label] == "section":
                    continue
                for pos in range(min(n_values, len(BUILDING_COLUMNS))):
                    j = pos + 1
                    raw = cell(cols[j], top, cols[j + 1], bottom, numeric=True)
                    value = parse_number(raw)
                    if value is None and raw == "":
                        # sel kosong pada baris nilai berarti nol pada formulir BTB
                        value = 0.0
                    sheet.values.loc[label, BUILDING_COLUMNS[pos]] = value

    if found_rows == 0:
        raise BTBError(
            "Tidak ada tabel BTB yang dikenali di dalam PDF ini. Pastikan berkas "
            "adalah hasil cetak halaman 'BTB Interaktif' dari portal MAPPI."
        )
    return sheet


# ---------------------------------------------------------------------------
# Pembacaan XLSX / CSV
# ---------------------------------------------------------------------------


def _read_any_table(file: Any, filename: str) -> pd.DataFrame:
    name = (filename or getattr(file, "name", "") or "").lower()
    if name.endswith((".csv", ".txt", ".tsv")):
        sep = "\t" if name.endswith(".tsv") else None
        return pd.read_csv(file, sep=sep, engine="python", dtype=object)
    excel = pd.ExcelFile(file, engine="openpyxl")
    sheet_name = SHEET_NAME if SHEET_NAME in excel.sheet_names else excel.sheet_names[0]
    return excel.parse(sheet_name, dtype=object)


def read_spreadsheet(file: Any, filename: str = "") -> list[BTBSheet]:
    """Baca satu atau banyak lembar BTB dari XLSX/CSV.

    Mendukung dua bentuk:

    * **lebar** — kolom identitas + 11 kolom tipe bangunan (bentuk yang ditulis
      aplikasi ini, dan bentuk paling dekat dengan dokumen aslinya);
    * **panjang** — kolom ``Tipe Bangunan`` dan ``Nilai (Rp/m2)``.
    """
    df = _read_any_table(file, filename)
    if df.empty:
        raise BTBError("Berkas tidak berisi data.")
    df.columns = [" ".join(str(c).split()) for c in df.columns]

    lookup = {_norm(c): c for c in df.columns}

    def col(*candidates: str) -> str | None:
        for cand in candidates:
            hit = lookup.get(_norm(cand))
            if hit:
                return hit
        return None

    elemen_col = col("Elemen", "Elemen Bangunan", "Uraian")
    if elemen_col is None:
        raise BTBError(
            "Kolom 'Elemen' tidak ditemukan. Gunakan berkas hasil unduhan "
            "aplikasi ini, atau sediakan kolom 'Elemen' beserta kolom tipe bangunan."
        )

    kota_col = col("Kota/Kabupaten", "Kota / Kabupaten", "Kota", "Kabupaten")
    prov_col = col("Provinsi")
    tahun_col = col("Tahun")
    tipe_col = col("Tipe Bangunan", "Tipe")
    nilai_col = col("Nilai (Rp/m2)", "Nilai", "Nilai (Rp/m²)")

    if tipe_col and nilai_col:  # bentuk panjang -> putar ke bentuk lebar
        df = df.copy()
        df["__tipe"] = df[tipe_col].map(lambda v: match_building_column(v) or str(v))
        index_cols = [c for c in (prov_col, kota_col, tahun_col, elemen_col) if c]
        df = df.pivot_table(
            index=index_cols, columns="__tipe", values=nilai_col, aggfunc="first"
        ).reset_index()
        lookup = {_norm(c): c for c in df.columns}

    # petakan kolom tipe bangunan yang ada di berkas
    mapping: dict[str, str] = {}
    for source in df.columns:
        if source in {elemen_col, kota_col, prov_col, tahun_col}:
            continue
        target = match_building_column(source)
        if target and target not in mapping.values():
            mapping[source] = target
    if not mapping:
        raise BTBError(
            "Tidak ada kolom tipe bangunan yang dikenali (Mewah, Menengah, "
            "Sederhana, Gudang, Low-Rise, Mid-Rise, High-Rise, Mall, Hotel, Apartemen)."
        )

    df["__elemen"] = df[elemen_col].map(match_row_label)
    df = df[df["__elemen"].notna()]
    if df.empty:
        raise BTBError("Tidak ada baris elemen BTB yang dikenali di dalam berkas.")

    df["__kota"] = df[kota_col].astype(str).str.strip() if kota_col else ""
    df["__prov"] = df[prov_col].astype(str).str.strip() if prov_col else ""
    if tahun_col:
        df["__tahun"] = df[tahun_col].map(lambda v: int(parse_number(v)) if parse_number(v) else None)
    else:
        df["__tahun"] = None

    sheets: list[BTBSheet] = []
    for (kota, tahun), part in df.groupby(["__kota", "__tahun"], dropna=False, sort=False):
        sheet = BTBSheet(
            provinsi=str(part["__prov"].iloc[0] or ""),
            kota_kabupaten=str(kota or ""),
            tahun=int(tahun) if pd.notna(tahun) and tahun is not None else None,
            source=filename or getattr(file, "name", "") or "Spreadsheet",
        )
        for _, row in part.iterrows():
            label = row["__elemen"]
            if ROW_KIND[label] == "section":
                continue
            for source_col, target in mapping.items():
                sheet.values.loc[label, target] = parse_number(row[source_col])
        missing = [c for c in BUILDING_COLUMNS if c not in mapping.values()]
        if missing:
            sheet.notes.append("Kolom tanpa data pada berkas: " + ", ".join(missing) + ".")
        sheets.append(sheet)
    return sheets


def read_any(file: Any, filename: str = "") -> list[BTBSheet]:
    """Baca berkas apa pun yang didukung (PDF/XLSX/XLSM/CSV)."""
    name = (filename or getattr(file, "name", "") or "").lower()
    if name.endswith(".pdf"):
        return [read_pdf(file, filename)]
    return read_spreadsheet(file, filename)


# ---------------------------------------------------------------------------
# Validasi aritmatika
# ---------------------------------------------------------------------------


def validate(sheet: BTBSheet, ppn_rate: float = DEFAULT_PPN_RATE, tolerance: float = 1.0) -> list[str]:
    """Periksa konsistensi angka; berguna untuk menangkap salah baca OCR."""
    issues: list[str] = []
    values = sheet.values

    def get(label: str, column: str) -> float | None:
        value = values.loc[label, column]
        return None if pd.isna(value) else float(value)

    for column in BUILDING_COLUMNS:
        kosong = [label for label in ROW_LABELS if ROW_KIND[label] != "section" and get(label, column) is None]
        if len(kosong) == len(ROW_LABELS) - 2:
            issues.append(f"{column}: seluruh nilai kosong.")
            continue
        if kosong:
            issues.append(f"{column}: {len(kosong)} nilai kosong ({', '.join(kosong[:3])}…).")

        checks: list[tuple[str, float | None, float | None]] = []
        items_a = [get(label, column) for label in ITEMS_A]
        items_b = [get(label, column) for label in ITEMS_B]
        total_a, total_b = get(TOTAL_A, column), get(TOTAL_B, column)
        total_ab, ppn = get(TOTAL_AB, column), get(PPN_ROW, column)
        total_ppn, rounded = get(TOTAL_PPN, column), get(ROUNDED, column)

        if all(v is not None for v in items_a):
            checks.append((TOTAL_A, total_a, sum(items_a)))
        if all(v is not None for v in items_b):
            checks.append((TOTAL_B, total_b, sum(items_b)))
        if total_a is not None and total_b is not None:
            checks.append((TOTAL_AB, total_ab, total_a + total_b))
        if total_ab is not None:
            checks.append((PPN_ROW, ppn, total_ab * ppn_rate))
            if ppn is not None:
                checks.append((TOTAL_PPN, total_ppn, total_ab + ppn))
        if total_ppn is not None:
            checks.append((ROUNDED, rounded, round(total_ppn, -4)))

        for label, actual, expected in checks:
            if actual is None or expected is None:
                continue
            tol = max(tolerance, abs(expected) * 5e-6)
            if label == ROUNDED:
                tol = 10_000  # pembulatan pada dokumen memakai kelipatan 10.000
            if abs(actual - expected) > tol:
                issues.append(
                    f"{column} – {label}: tertulis {actual:,.0f}, "
                    f"perhitungan {expected:,.0f} (selisih {actual - expected:,.0f})."
                )
    return issues


def recompute_totals(sheet: BTBSheet, ppn_rate: float = DEFAULT_PPN_RATE) -> BTBSheet:
    """Hitung ulang seluruh baris total dari komponen biayanya."""
    values = sheet.values
    for column in BUILDING_COLUMNS:
        items_a = [values.loc[label, column] for label in ITEMS_A]
        items_b = [values.loc[label, column] for label in ITEMS_B]
        if any(pd.isna(v) for v in items_a + items_b):
            continue
        total_a = float(sum(float(v) for v in items_a))
        total_b = float(sum(float(v) for v in items_b))
        total_ab = total_a + total_b
        ppn = round(total_ab * ppn_rate)
        values.loc[TOTAL_A, column] = total_a
        values.loc[TOTAL_B, column] = total_b
        values.loc[TOTAL_AB, column] = total_ab
        values.loc[PPN_ROW, column] = ppn
        values.loc[TOTAL_PPN, column] = total_ab + ppn
        values.loc[ROUNDED, column] = round(total_ab + ppn, -4)
    return sheet


# ---------------------------------------------------------------------------
# Penulisan XLSX
# ---------------------------------------------------------------------------

_NUMBER_FORMAT = "#,##0"


def _style_worksheet(worksheet, df: pd.DataFrame, freeze: str = "A2") -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    numeric_cols = {
        i + 1
        for i, name in enumerate(df.columns)
        if name in BUILDING_COLUMNS or name in {"Nilai (Rp/m2)", "Tahun", "No"}
    }
    bold_labels = {TOTAL_A, TOTAL_B, TOTAL_AB, TOTAL_PPN, ROUNDED, PPN_ROW}
    label_positions = [i + 1 for i, name in enumerate(df.columns) if name == "Elemen"]

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        is_total = any(row[pos - 1].value in bold_labels for pos in label_positions)
        for cell in row:
            cell.border = border
            if cell.column in numeric_cols:
                cell.alignment = Alignment(horizontal="right")
                if df.columns[cell.column - 1] not in {"Tahun", "No"}:
                    cell.number_format = _NUMBER_FORMAT
            if is_total:
                cell.font = Font(bold=True)

    for i, name in enumerate(df.columns, start=1):
        longest = max([len(str(name))] + [len(f"{v:,.0f}") if isinstance(v, (int, float)) and not pd.isna(v) else len(str(v)) for v in df[name].head(400)])
        worksheet.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 38)
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = worksheet.dimensions


def write_workbook(path_or_buffer: Any, wide: pd.DataFrame, include_long: bool = True) -> Any:
    """Tulis data BTB ke XLSX yang siap pakai di Excel."""
    wide = wide.copy()
    for column in BUILDING_COLUMNS:
        if column in wide.columns:
            wide[column] = pd.to_numeric(wide[column], errors="coerce")
    if "Tahun" in wide.columns:
        wide["Tahun"] = pd.to_numeric(wide["Tahun"], errors="coerce").astype("Int64")

    with pd.ExcelWriter(path_or_buffer, engine="openpyxl") as writer:
        wide.to_excel(writer, index=False, sheet_name=SHEET_NAME)
        _style_worksheet(writer.sheets[SHEET_NAME], wide)
        if include_long:
            long = wide_to_long(wide)
            long.to_excel(writer, index=False, sheet_name="BTB Long")
            _style_worksheet(writer.sheets["BTB Long"], long)
    return path_or_buffer


def wide_to_long(wide: pd.DataFrame) -> pd.DataFrame:
    id_cols = [c for c in ID_COLUMNS if c in wide.columns]
    value_cols = [c for c in BUILDING_COLUMNS if c in wide.columns]
    long = wide.melt(id_vars=id_cols, value_vars=value_cols, var_name="Tipe Bangunan", value_name="Nilai (Rp/m2)")
    spec = {name: (grp, sp) for grp, name, sp, _ in COLUMN_SPECS}
    long["Kelompok Bangunan"] = long["Tipe Bangunan"].map(lambda c: spec.get(c, ("", ""))[0])
    long["Spesifikasi"] = long["Tipe Bangunan"].map(lambda c: spec.get(c, ("", ""))[1])
    order = id_cols + ["Kelompok Bangunan", "Tipe Bangunan", "Spesifikasi", "Nilai (Rp/m2)"]
    sort_cols = [c for c in ("Kota/Kabupaten", "Tahun", "No") if c in long.columns]
    return long[order].sort_values(sort_cols + ["Tipe Bangunan"], kind="stable").reset_index(drop=True)


def load_workbook_data(path: str) -> pd.DataFrame:
    """Baca kembali basis data XLSX aplikasi. Kosong bila berkas belum ada."""
    if not os.path.exists(path):
        return pd.DataFrame(columns=ID_COLUMNS + BUILDING_COLUMNS)
    excel = pd.ExcelFile(path, engine="openpyxl")
    sheet_name = SHEET_NAME if SHEET_NAME in excel.sheet_names else excel.sheet_names[0]
    df = excel.parse(sheet_name)
    df.columns = [" ".join(str(c).split()) for c in df.columns]
    if "Tahun" in df.columns:
        df["Tahun"] = pd.to_numeric(df["Tahun"], errors="coerce").astype("Int64")
    for column in BUILDING_COLUMNS:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")
    return df


def save_sheet(path: str, sheet: BTBSheet, replace_existing: bool = True) -> pd.DataFrame:
    """Simpan satu lembar BTB ke basis data XLSX dan kembalikan isi terbarunya.

    Menyimpan ulang kota/tahun yang sama akan **menimpa** data lama
    (bukan menambah duplikat) selama ``replace_existing`` bernilai benar.
    """
    if not sheet.is_complete():
        raise BTBError("Kota/Kabupaten dan Tahun wajib diisi sebelum menyimpan.")

    existing = load_workbook_data(path)
    incoming = sheet.to_wide()

    if not existing.empty and replace_existing:
        mask = (
            existing["Kota/Kabupaten"].astype(str).str.strip().str.casefold()
            == sheet.kota_kabupaten.strip().casefold()
        ) & (pd.to_numeric(existing["Tahun"], errors="coerce") == sheet.tahun)
        existing = existing[~mask]

    combined = pd.concat([existing, incoming], ignore_index=True) if not existing.empty else incoming
    for column in ID_COLUMNS + BUILDING_COLUMNS:
        if column not in combined.columns:
            combined[column] = pd.NA
    combined = combined[ID_COLUMNS + BUILDING_COLUMNS]
    combined = combined.sort_values(["Kota/Kabupaten", "Tahun", "No"], kind="stable").reset_index(drop=True)

    directory = os.path.dirname(os.path.abspath(path))
    os.makedirs(directory, exist_ok=True)
    write_workbook(path, combined)
    return combined


def build_download(wide: pd.DataFrame, layout: str = "wide") -> io.BytesIO:
    """Bangun berkas XLSX untuk diunduh (``wide`` atau ``long``)."""
    buffer = io.BytesIO()
    if layout == "long":
        long = wide_to_long(wide)
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            long.to_excel(writer, index=False, sheet_name="BTB Long")
            _style_worksheet(writer.sheets["BTB Long"], long)
    else:
        write_workbook(buffer, wide, include_long=False)
    buffer.seek(0)
    return buffer


def summary(df: pd.DataFrame) -> pd.DataFrame:
    """Ringkasan kota/kabupaten dan tahun yang sudah tersimpan."""
    if df.empty:
        return pd.DataFrame(columns=["Provinsi", "Kota/Kabupaten", "Tahun", "Jumlah Baris"])
    grouped = (
        df.groupby([c for c in ("Provinsi", "Kota/Kabupaten", "Tahun") if c in df.columns], dropna=False)
        .size()
        .reset_index(name="Jumlah Baris")
    )
    return grouped.sort_values(["Kota/Kabupaten", "Tahun"], kind="stable").reset_index(drop=True)
